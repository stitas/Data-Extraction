from google.api_core.client_options import ClientOptions
from google.cloud.documentai_toolbox import document
from google.cloud import documentai, storage
from dotenv import load_dotenv
from io import BytesIO
from PIL import Image
import pandas as pd
import openpyxl
import datetime
import base64
import json
import fitz
import xlrd

load_dotenv('.env')

# project_id = 'dataextraction-403815'
# location = 'eu'
# text_parser_processor_id = 'c2a06540e73d5905'
# form_parser_processor_id = 'c7ae90b5bf93bf21'
# gcs_input_prefix = 'gs://sample_extractor_files/'
# bucket_name = 'sample_extractor_files'
# extracted_tables_bucket_name = 'extracted_tables' 

MIME_TYPES = {
    'pdf':	'application/pdf',
    'gif':	'image/gif',
    'tiff': 'image/tiff',
    'tif': 'image/tiff',
    'jpg': 'image/jpeg',
    'jpeg': 'image/jpeg',
    'png': 'image/png',
    'bmp': 'image/bmp',
    'webp': 'image/webp'
}


def get_mime_type(gcs_uri):
    return MIME_TYPES[gcs_uri.rsplit('.',1)[1]]    

def get_file_name_from_gcs_uri(gcs_uri):
    return gcs_uri.rsplit('/',1)[1]

def get_file_bytes(content):
    return content.rsplit(',',1)[1]
    
# Analyze document   
def process_document(project_id, location, processor_id, gcs_uri, mime_type):
    opts = ClientOptions(api_endpoint=f'{location}-documentai.googleapis.com')
    client = documentai.DocumentProcessorServiceClient(client_options=opts)
    name = client.processor_path(project_id, location, processor_id)
    
    # Optional: Additional configurations for processing.
    # For more information: https://cloud.google.com/document-ai/docs/reference/rest/v1/ProcessOptions
    
    gcs_document = documentai.GcsDocument(gcs_uri=gcs_uri, mime_type=mime_type)

    request = documentai.ProcessRequest(name=name, gcs_document=gcs_document)
    result = client.process_document(request=request)
    document = result.document
    
    return document
 
def process_document_binary(project_id, location, processor_id, file_content, mime_type):
    opts = ClientOptions(api_endpoint=f"{location}-documentai.googleapis.com")
    client = documentai.DocumentProcessorServiceClient(client_options=opts)
    name = client.processor_path(project_id, location, processor_id)

    # Load binary data
    raw_document = documentai.RawDocument(content=file_content, mime_type=mime_type)

    # Configure the process request
    request = documentai.ProcessRequest(name=name, raw_document=raw_document)
    result = client.process_document(request=request)
    document = result.document

    return document
 
# Upload document to bucket   
def upload_blob(bucket_name, file_obj, destination_blob_name):
    storage_client = storage.Client()
    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(destination_blob_name)
    
    blob.upload_from_file(file_obj)
   
# Delete document from bucket 
def delete_blob(bucket_name, blob_name):
    storage_client = storage.Client()
    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(blob_name)
    
    blob.delete()
    
def get_kvp(project_id, location, processor_id, gcs_uri, mime_type):
    document = process_document(project_id, location, processor_id, gcs_uri, mime_type)
    
    kvp = []
    pair = {
        'key': '',
        'value': '',
        'key_confidence': 0,
        'value_confidence': 0,
    }
    
    for page in document.pages:
        for i in range(len(page.form_fields)):
            pair['key'] = page.form_fields[i].field_name.text_anchor.content
            pair['value'] = page.form_fields[i].field_value.text_anchor.content
            pair['key_confidence'] = page.form_fields[i].field_name.confidence * 100
            pair['value_confidence'] = page.form_fields[i].field_value.confidence * 100
            kvp.append(pair)
            pair = {
                'key': '',
                'value': '',
                'key_confidence': 0,
                'value_confidence': 0,
            }
        
    return kvp

def upload_table(project_id, location, processor_id, gcs_uri, mime_type, extracted_tables_bucket_name):
    proccessed_document = process_document(project_id, location, processor_id, gcs_uri, mime_type)
    wrapped_document = document.Document.from_documentai_document(proccessed_document)

    paths = []
    vertices_list = []
    
    vertices = {
        'x1': 0,
        'y1': 0,
        'x2': 0,
        'y2': 0,
    }

    for page in proccessed_document.pages:
        for i in range(len(page.tables)):
            vertices['x1'] = page.tables[i].layout.bounding_poly.normalized_vertices[0].x
            vertices['y1'] = page.tables[i].layout.bounding_poly.normalized_vertices[0].y
            vertices['x2'] = page.tables[i].layout.bounding_poly.normalized_vertices[2].x
            vertices['y2'] = page.tables[i].layout.bounding_poly.normalized_vertices[2].y
            vertices_list.append(vertices)
            vertices = {
                'x1': 0,
                'y1': 0,
                'x2': 0,
                'y2': 0,
            }

    for page in wrapped_document.pages:
        for i in range(len(page.tables)):
            df = page.tables[i].to_dataframe()
            path = 'gs://'+extracted_tables_bucket_name+'/'+datetime.datetime.now().strftime('%Y-%m-%d_%H:%M:%S:%f')+'.xlsx'
            df.to_excel(path)
            paths.append(path)            

    return paths, vertices_list

def get_dataframes(paths):
    dataframes = []
    
    for path in paths:
        df = pd.read_excel(path)
        df = df.dropna()
        for i in df.columns:
            if 'Unnamed' in i:
                df.drop(i, inplace=True, axis=1)
            
        dataframes.append(df)
        
    return dataframes
    
def stream_pdf_page_as_image(bucket_name, file_name, page_number):
    storage_client = storage.Client()
    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(file_name)
    matrix = fitz.Matrix(2,2)
    
    pdf_content = BytesIO()
    blob.download_to_file(pdf_content)
    
    pdf_document = fitz.open("pdf", pdf_content.getvalue())
    page = pdf_document[page_number]
    image_bytes = page.get_pixmap(matrix=matrix).tobytes()
    
    image = base64.b64encode(image_bytes).decode('utf-8')
    
    return image

def get_pdf_page_count_gcs(bucket_name, file_name):
    storage_client = storage.Client()
    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(file_name)

    pdf_content = BytesIO()
    blob.download_to_file(pdf_content)

    pdf_document = fitz.open('pdf', pdf_content.getvalue())

    page_count = pdf_document.page_count

    pdf_document.close()

    return page_count

def get_pdf_page_count(file_content):
    doc = fitz.open('pdf', file_content)
    page_count = doc.page_count
    
    return page_count

def get_text_from_verticies(file_content, key, x1, y1, x2, y2, project_id, location, processor_id, mime_type, page):
    left_top = fitz.Point(x1, y1)
    right_bot = fitz.Point(x2, y2)
    zoom = 2 # to increase the resolution
    mat = fitz.Matrix(zoom, zoom)

    doc = fitz.open('pdf', file_content)

    rect = fitz.Rect(left_top, right_bot)
    doc[page].set_cropbox(rect)

    pix = doc[0].get_pixmap(matrix=mat)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

    buff = BytesIO()
    img.save(buff, format="JPEG")
    img_str = buff.getvalue()
    
    text = process_document_binary(project_id, location, processor_id, img_str, mime_type).text
    
    return {'key':key, 'value':text}
    

def get_kvp_from_key(document, key):
    kvps = []
    for page in document.pages:
        for field in page.form_fields:
            if field.field_name.text_anchor.content == key:
                key = field.field_name.text_anchor.content
                value = field.field_value.text_anchor.content
                kvps.append({'key':key, 'value':value})
        
    return kvps

def get_table_from_vertices(proccessed_document, x1, y1, x2, y2):
    df = pd.DataFrame()
    for i in range(len(proccessed_document.pages)):
        for j in range(len(proccessed_document.pages[i].tables)):
            if proccessed_document.pages[i].tables[j].layout.bounding_poly.normalized_vertices[0].x == x1 and proccessed_document.pages[i].tables[j].layout.bounding_poly.normalized_vertices[0].y == y1 and proccessed_document.pages[i].tables[j].layout.bounding_poly.normalized_vertices[2].x == x2 and proccessed_document.pages[i].tables[j].layout.bounding_poly.normalized_vertices[2].y == y2:
                wrapped_document = document.Document.from_documentai_document(proccessed_document)
                df = wrapped_document.pages[i].tables[j].to_dataframe()
                
    
    return {'table_as_json':df.to_json()}